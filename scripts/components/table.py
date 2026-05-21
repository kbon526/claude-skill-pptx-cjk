"""轻量风格表格组件。

视觉:浅灰表头 + 横线分隔 + 红色底边(表头) + 无外边框。
设计目标:中文商务提案标准表格风格,避免 PowerPoint 默认蓝绿底色俗气感。
"""

from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

from ..core.font import _sf, SZ_CAP
from ..core.color import DARK_GRAY, LIGHT_GRAY


def _set_cell_border(cell, side, width_pt, color_hex):
    """设置单元格某一侧边框。width_pt=0 表示显式无边框。"""
    tc = cell._tc
    tcPr = tc.tcPr
    if tcPr is None:
        tcPr = etree.SubElement(tc, qn("a:tcPr"))
    tag = {"left": "a:lnL", "right": "a:lnR", "top": "a:lnT", "bottom": "a:lnB"}[side]
    ln = tcPr.find(qn(tag))
    if ln is not None:
        tcPr.remove(ln)
    if width_pt <= 0:
        ln = etree.SubElement(tcPr, qn(tag), attrib={"w": "0"})
        etree.SubElement(ln, qn("a:noFill"))
        return
    ln = etree.SubElement(tcPr, qn(tag), attrib={"w": str(int(width_pt * 12700))})
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), attrib={"val": color_hex})


def _style_cell(cell, text, sz=SZ_CAP, bold=False, color=None, fill=None,
                align=PP_ALIGN.CENTER):
    """统一注入字体 + 填充 + 对齐。"""
    cell.text = str(text)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            _sf(r, sz, bold, color)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()


def add_table(slide, l, t, w, h, data, col_widths=None, first_col_bold=True,
              accent_hex="E60000"):
    """添加轻风格表格。

    Args:
        data: 二维列表,第一行为表头
        col_widths: 列宽列表(Inches),不传则等宽
        first_col_bold: 第一列是否加粗(列标签场景)
        accent_hex: 表头底边强调色,默认红
    """
    rows = len(data)
    cols = len(data[0]) if data else 0
    ts = slide.shapes.add_table(rows, cols, l, t, w, h)
    tbl = ts.table

    # 关闭默认条带样式
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("bandCol", "0")
    tblPr.set("firstRow", "0")
    tblPr.set("firstCol", "0")

    # 列宽
    if col_widths:
        for i, cw in enumerate(col_widths):
            if i < cols:
                tbl.columns[i].width = cw
    else:
        cw = int(w / cols)
        for i in range(cols):
            tbl.columns[i].width = cw

    # 单元格内容 + 边框
    for ri in range(rows):
        for ci in range(cols):
            cell = tbl.cell(ri, ci)
            val = data[ri][ci] if ci < len(data[ri]) else ""
            if ri == 0:
                _style_cell(cell, val, SZ_CAP, True, DARK_GRAY, LIGHT_GRAY)
            elif first_col_bold and ci == 0:
                _style_cell(cell, val, SZ_CAP, True, DARK_GRAY, None,
                            align=PP_ALIGN.LEFT)
            else:
                _style_cell(cell, val, SZ_CAP, False, DARK_GRAY, None)
            # 边框规则:无左右,横线分隔,表头底部红
            _set_cell_border(cell, "left", 0, "FFFFFF")
            _set_cell_border(cell, "right", 0, "FFFFFF")
            if ri == 0:
                _set_cell_border(cell, "top", 0, "FFFFFF")
                _set_cell_border(cell, "bottom", 1.5, accent_hex)
            else:
                _set_cell_border(cell, "top", 0.5, "D9D9D9")
                if ri == rows - 1:
                    _set_cell_border(cell, "bottom", 0.75, "999999")
                else:
                    _set_cell_border(cell, "bottom", 0, "FFFFFF")
    return ts
